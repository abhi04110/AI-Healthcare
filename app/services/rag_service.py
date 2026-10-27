
"""Medical knowledge RAG ingestion and retrieval.

Supports PDF, TXT, Markdown, DOCX, XLSX, XLS, CSV and common image files.
Patient-specific records must remain in SQL and follow existing access controls.
"""

from __future__ import annotations

import csv
import hashlib
import io
import os
import re
import zipfile
from functools import lru_cache
from pathlib import Path
from typing import Any

from app.config import (
    MEDICAL_KNOWLEDGE_DIR,
    RAG_CHUNK_OVERLAP,
    RAG_CHUNK_SIZE,
    RAG_EMBEDDING_MODEL,
    RAG_TOP_K,
    VECTOR_STORE_DIR,
)

SUPPORTED_EXTENSIONS = {
    ".pdf", ".txt", ".md", ".markdown", ".docx", ".xlsx", ".xls",
    ".csv", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
MAX_EMBEDDED_IMAGE_BYTES = 20 * 1024 * 1024
COLLECTION_NAME = "medical_knowledge"


def _dependencies():
    try:
        import chromadb
        from sentence_transformers import SentenceTransformer
    except ImportError as exc:
        raise RuntimeError(
            "RAG dependencies missing. Install requirements.txt first."
        ) from exc
    return chromadb, SentenceTransformer


@lru_cache(maxsize=1)
def _embedding_model():
    _, model_class = _dependencies()
    return model_class(RAG_EMBEDDING_MODEL)


@lru_cache(maxsize=1)
def _collection():
    chromadb, _ = _dependencies()
    Path(VECTOR_STORE_DIR).mkdir(parents=True, exist_ok=True)
    client = chromadb.PersistentClient(path=str(VECTOR_STORE_DIR))
    return client.get_or_create_collection(
        name=COLLECTION_NAME,
        metadata={"hnsw:space": "cosine"},
    )


def _source_key(relative_path: str) -> str:
    return hashlib.sha256(relative_path.encode("utf-8")).hexdigest()[:24]


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _ocr_bytes(data: bytes, filename: str = "image") -> str:
    try:
        import pytesseract
        from PIL import Image, ImageEnhance, ImageFilter, ImageOps

        windows_tesseract = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
        if os.path.exists(windows_tesseract):
            pytesseract.pytesseract.tesseract_cmd = windows_tesseract

        image = Image.open(io.BytesIO(data)).convert("RGB")
        image = ImageOps.exif_transpose(image)

        if image.width < 1600:
            ratio = min(2.0, 1600 / max(1, image.width))
            image = image.resize(
                (int(image.width * ratio), int(image.height * ratio))
            )

        image = ImageEnhance.Contrast(image).enhance(1.25)
        image = image.filter(ImageFilter.SHARPEN)
        return pytesseract.image_to_string(image, config="--psm 3").strip()
    except Exception as exc:
        raise RuntimeError(
            f"OCR failed for '{filename}': {type(exc).__name__}: {exc}. "
            "Check Pillow, pytesseract and Tesseract OCR installation."
        ) from exc


def _unique_text(native: str, ocr: str) -> str:
    native = native.strip()
    normalized_native = re.sub(r"[^a-z0-9]+", " ", native.lower())
    additions = []
    seen = set()

    for line in ocr.splitlines():
        line = line.strip()
        normalized = re.sub(r"[^a-z0-9]+", " ", line.lower()).strip()

        if len(normalized) < 3 or normalized in seen:
            continue

        seen.add(normalized)

        if normalized not in normalized_native:
            additions.append(line)

    if additions:
        return (
            native
            + "\n\n[OCR text from scanned or embedded content]\n"
            + "\n".join(additions)
        ).strip()

    return native


def _read_pdf(path: Path) -> list[tuple[int, str]]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    results = []

    for page_number, page in enumerate(reader.pages, start=1):
        native_text = (page.extract_text() or "").strip()
        combined = native_text

        try:
            import fitz

            with fitz.open(str(path)) as document:
                pdf_page = document[page_number - 1]

                if len(native_text) < 80:
                    pix = pdf_page.get_pixmap(
                        matrix=fitz.Matrix(1.7, 1.7),
                        alpha=False,
                    )
                    ocr_text = _ocr_bytes(
                        pix.tobytes("png"),
                        f"{path.name}, page {page_number}",
                    )
                    combined = _unique_text(combined, ocr_text)

                for block in pdf_page.get_text("dict").get("blocks", []):
                    if block.get("type") != 1:
                        continue

                    bbox = block.get("bbox")
                    if not bbox:
                        continue

                    pix = pdf_page.get_pixmap(
                        matrix=fitz.Matrix(2, 2),
                        clip=fitz.Rect(bbox),
                        alpha=False,
                    )

                    if len(pix.samples) <= MAX_EMBEDDED_IMAGE_BYTES:
                        ocr_text = _ocr_bytes(
                            pix.tobytes("png"),
                            f"{path.name}, page {page_number} image",
                        )
                        combined = _unique_text(combined, ocr_text)

        except ImportError:
            if not native_text:
                raise RuntimeError(
                    "Scanned PDF OCR requires PyMuPDF. Install pymupdf."
                )

        results.append((page_number, combined))

    return results


def _read_docx(path: Path) -> list[tuple[int, str]]:
    from docx import Document

    document = Document(str(path))
    pieces = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    for table_number, table in enumerate(document.tables, start=1):
        pieces.append(f"Table {table_number}:")
        for row in table.rows:
            pieces.append(
                " | ".join(
                    cell.text.strip().replace("\n", " / ")
                    for cell in row.cells
                )
            )

    try:
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                if not name.startswith("word/media/") or name.endswith("/"):
                    continue

                data = archive.read(name)
                if len(data) > MAX_EMBEDDED_IMAGE_BYTES:
                    continue

                try:
                    text = _ocr_bytes(data, f"{path.name}:{Path(name).name}")
                    if text:
                        pieces.append(f"[Embedded image OCR]\n{text}")
                except Exception:
                    # Do not silently pretend image text was extracted.
                    pieces.append(
                        f"[Warning: OCR failed for embedded image {Path(name).name}]"
                    )
    except zipfile.BadZipFile as exc:
        raise ValueError(f"Invalid DOCX file: {path.name}") from exc

    return [(1, "\n".join(pieces))]


def _read_excel(path: Path) -> list[tuple[int, str]]:
    import pandas as pd

    sheets = pd.read_excel(path, sheet_name=None, dtype=str)
    results = []

    for sheet_number, (sheet_name, frame) in enumerate(sheets.items(), start=1):
        frame = frame.fillna("")
        lines = [f"Workbook: {path.name}\nSheet: {sheet_name}"]

        if not frame.empty:
            lines.append(" | ".join(str(column) for column in frame.columns))

            for row in frame.itertuples(index=False, name=None):
                values = [str(value).strip() for value in row]
                if any(values):
                    lines.append(" | ".join(values))

        results.append((sheet_number, "\n".join(lines)))

    if path.suffix.lower() == ".xlsx":
        try:
            with zipfile.ZipFile(path) as archive:
                for name in archive.namelist():
                    if not name.startswith("xl/media/") or name.endswith("/"):
                        continue

                    data = archive.read(name)
                    if len(data) > MAX_EMBEDDED_IMAGE_BYTES:
                        continue

                    try:
                        text = _ocr_bytes(data, f"{path.name}:{Path(name).name}")
                        if text:
                            results.append(
                                (
                                    len(results) + 1,
                                    f"Workbook: {path.name}\n{text}",
                                )
                            )
                    except Exception:
                        results.append(
                            (
                                len(results) + 1,
                                f"Workbook: {path.name}\n"
                                f"[Warning: OCR failed for {Path(name).name}]",
                            )
                        )
        except zipfile.BadZipFile as exc:
            raise ValueError(f"Invalid XLSX file: {path.name}") from exc

    return results


def _read_csv(path: Path) -> list[tuple[int, str]]:
    try:
        content = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError:
        content = path.read_text(encoding="cp1252", errors="replace")

    try:
        dialect = csv.Sniffer().sniff(content[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel

    rows = csv.reader(io.StringIO(content), dialect)
    text = "\n".join(
        " | ".join(cell.strip() for cell in row)
        for row in rows
        if any(cell.strip() for cell in row)
    )
    return [(1, text)]


def _read_document(path: Path) -> list[tuple[int, str]]:
    extension = path.suffix.lower()

    if extension == ".pdf":
        return _read_pdf(path)
    if extension == ".docx":
        return _read_docx(path)
    if extension in {".xlsx", ".xls"}:
        return _read_excel(path)
    if extension == ".csv":
        return _read_csv(path)
    if extension in IMAGE_EXTENSIONS:
        return [(1, _ocr_bytes(path.read_bytes(), path.name))]

    return [(1, path.read_text(encoding="utf-8-sig", errors="replace"))]


def _clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _split_text(text: str) -> list[str]:
    text = _clean_text(text)
    if not text:
        return []

    size = max(200, int(RAG_CHUNK_SIZE))
    overlap = max(0, min(int(RAG_CHUNK_OVERLAP), size // 2))
    chunks = []
    start = 0

    while start < len(text):
        end = min(start + size, len(text))

        if end < len(text):
            boundaries = [
                text.rfind("\n\n", start, end),
                text.rfind(". ", start, end),
                text.rfind("? ", start, end),
                text.rfind("! ", start, end),
            ]
            boundary = max(boundaries)

            if boundary > start + size // 2:
                end = boundary + 1

        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)

        if end >= len(text):
            break

        start = max(start + 1, end - overlap)

    return chunks


def _existing_source_hash(collection, source_key: str) -> str | None:
    result = collection.get(
        where={"source_key": source_key},
        include=["metadatas"],
    )
    metadatas = result.get("metadatas") or []

    for metadata in metadatas:
        if metadata and metadata.get("source_hash"):
            return str(metadata["source_hash"])

    return None


def _delete_source_chunks(collection, source_key: str) -> int:
    result = collection.get(
        where={"source_key": source_key},
        include=["metadatas"],
    )
    ids = result.get("ids") or []

    if ids:
        collection.delete(ids=ids)

    return len(ids)


def _index_file(path: Path, relative_source: str, collection, model) -> dict[str, Any]:
    source_key = _source_key(relative_source)
    source_hash = _file_hash(path)
    existing_hash = _existing_source_hash(collection, source_key)

    if existing_hash == source_hash:
        return {
            "source": relative_source,
            "status": "unchanged",
            "chunks": 0,
        }

    chunks_with_pages = []

    for page_number, page_text in _read_document(path):
        chunks_with_pages.extend(
            (page_number, chunk)
            for chunk in _split_text(page_text)
        )

    if not chunks_with_pages:
        return {
            "source": relative_source,
            "status": "empty_text",
            "chunks": 0,
            "warning": "No extractable text found.",
        }

    texts = [chunk for _, chunk in chunks_with_pages]
    vectors = model.encode(texts, normalize_embeddings=True).tolist()

    ids = [
        hashlib.sha256(
            f"{source_key}:{source_hash}:{index}".encode("utf-8")
        ).hexdigest()
        for index in range(len(texts))
    ]

    metadatas = [
        {
            "source": relative_source,
            "source_key": source_key,
            "source_hash": source_hash,
            "page": int(page_number),
            "chunk_index": int(index),
        }
        for index, (page_number, _) in enumerate(chunks_with_pages)
    ]

    # Build embeddings first. Replace the old version only after extraction
    # and embedding generation succeeded.
    _delete_source_chunks(collection, source_key)

    for offset in range(0, len(texts), 100):
        collection.add(
            ids=ids[offset:offset + 100],
            documents=texts[offset:offset + 100],
            embeddings=vectors[offset:offset + 100],
            metadatas=metadatas[offset:offset + 100],
        )

    return {
        "source": relative_source,
        "status": "indexed" if existing_hash is None else "updated",
        "chunks": len(texts),
    }


def ingest_medical_documents() -> dict[str, Any]:
    """Incrementally index supported documents and remove stale deleted sources."""
    root = Path(MEDICAL_KNOWLEDGE_DIR).resolve()
    root.mkdir(parents=True, exist_ok=True)

    files = sorted(
        path for path in root.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
    )

    collection = _collection()
    model = _embedding_model()
    results = []
    successful_source_keys = set()

    for path in files:
        relative_source = path.relative_to(root).as_posix()
        key = _source_key(relative_source)

        try:
            result = _index_file(path, relative_source, collection, model)
            results.append(result)

            if result["status"] in {"indexed", "updated", "unchanged"}:
                successful_source_keys.add(key)

        except Exception as exc:
            results.append({
                "source": relative_source,
                "status": "error",
                "chunks": 0,
                "error": f"{type(exc).__name__}: {exc}",
            })

    # Only remove sources confirmed absent from disk. If a file exists but
    # failed to process, retain its previous indexed version.
    current_paths = {
        _source_key(path.relative_to(root).as_posix())
        for path in files
    }

    indexed = collection.get(include=["metadatas"])
    metadata_list = indexed.get("metadatas") or []
    stale_keys = set()

    for metadata in metadata_list:
        if not metadata:
            continue

        key = metadata.get("source_key")
        if key and key not in current_paths:
            stale_keys.add(str(key))

    deleted_chunks = 0
    for key in stale_keys:
        deleted_chunks += _delete_source_chunks(collection, key)

    failed = sum(item["status"] == "error" for item in results)
    indexed_count = sum(
        item["chunks"]
        for item in results
        if item["status"] in {"indexed", "updated"}
    )

    return {
        "status": "partial_success" if failed else "success",
        "documents_found": len(files),
        "chunks_added_or_updated": indexed_count,
        "unchanged_files_skipped": sum(
            item["status"] == "unchanged" for item in results
        ),
        "deleted_sources_cleaned": len(stale_keys),
        "deleted_chunks": deleted_chunks,
        "failed_files": failed,
        "files": results,
        "vector_store": str(VECTOR_STORE_DIR),
        "embedding_model": RAG_EMBEDDING_MODEL,
    }


def retrieve_medical_knowledge(
    question: str,
    top_k: int | None = None,
) -> list[dict[str, Any]]:
    """Retrieve only the most relevant indexed passages for a question."""
    question = question.strip()
    if not question:
        return []

    collection = _collection()
    if collection.count() == 0:
        return []

    model = _embedding_model()
    query_vector = model.encode(
        [question],
        normalize_embeddings=True,
    ).tolist()[0]

    limit = max(1, min(int(top_k or RAG_TOP_K), 8))
    result = collection.query(
        query_embeddings=[query_vector],
        n_results=min(limit, collection.count()),
        include=["documents", "metadatas", "distances"],
    )

    documents = (result.get("documents") or [[]])[0]
    metadatas = (result.get("metadatas") or [[]])[0]
    distances = (result.get("distances") or [[]])[0]
    retrieved = []

    for text, metadata, distance in zip(documents, metadatas, distances):
        metadata = metadata or {}
        similarity = max(0.0, min(1.0, 1.0 - float(distance)))

        retrieved.append({
            "text": text,
            "source": metadata.get("source", "Unknown source"),
            "page": metadata.get("page"),
            "similarity": round(similarity, 4),
        })

    return retrieved


def format_medical_knowledge_context(
    chunks: list[dict[str, Any]],
) -> str:
    if not chunks:
        return "No relevant indexed medical reference passages were found."

    sections = []

    for index, chunk in enumerate(chunks, start=1):
        page = (
            f", page/sheet {chunk['page']}"
            if chunk.get("page") is not None
            else ""
        )
        sections.append(
            f"[Source {index}: {chunk['source']}{page}]\n{chunk['text']}"
        )

    return "\n\n".join(sections)


def list_medical_documents() -> list[dict[str, Any]]:
    """List source files and the number of indexed chunks for each file."""
    root = Path(MEDICAL_KNOWLEDGE_DIR).resolve()
    root.mkdir(parents=True, exist_ok=True)
    collection = _collection()
    documents = []

    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue

        relative_source = path.relative_to(root).as_posix()
        key = _source_key(relative_source)
        result = collection.get(
            where={"source_key": key},
            include=["metadatas"],
        )
        chunk_count = len(result.get("ids") or [])

        documents.append({
            "filename": relative_source,
            "size_bytes": path.stat().st_size,
            "indexed_chunks": chunk_count,
            "indexed": chunk_count > 0,
        })

    return documents


def delete_medical_document(filename: str) -> dict[str, Any]:
    """Delete a source file and its indexed chunks."""
    root = Path(MEDICAL_KNOWLEDGE_DIR).resolve()
    requested = Path(filename)

    if requested.is_absolute() or ".." in requested.parts:
        raise ValueError("Filename must be a relative path from the knowledge list.")

    target = (root / requested).resolve()

    if root not in target.parents:
        raise ValueError("Filename must stay inside the medical knowledge directory.")

    if not target.is_file() or target.suffix.lower() not in SUPPORTED_EXTENSIONS:
        return {
            "deleted": False,
            "filename": filename,
            "indexed_chunks_deleted": 0,
        }

    relative_source = target.relative_to(root).as_posix()
    key = _source_key(relative_source)
    deleted_chunks = _delete_source_chunks(_collection(), key)
    target.unlink()

    return {
        "deleted": True,
        "filename": relative_source,
        "indexed_chunks_deleted": deleted_chunks,
    }
